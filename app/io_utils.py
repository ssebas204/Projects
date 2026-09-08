"""File adapters. Imported Excel data is read only; exports are separate files."""
import io
import json
import unicodedata
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


def normalize(value):
    return "".join(c for c in unicodedata.normalize("NFD", str(value or "").lower()) if unicodedata.category(c) != "Mn").strip()


def code(value):
    if value is None:
        raise ValueError("Hay un código vacío en los archivos.")
    return str(int(value)) if isinstance(value, (int, float)) and float(value).is_integer() else str(value).strip()


def import_three(assignment, demand_file, matrix_file, base_config):
    def rows(source):
        return list(openpyxl.load_workbook(io.BytesIO(source), data_only=True, read_only=True).active.values)
    a, d, m = rows(assignment), rows(demand_file), rows(matrix_file)
    if len(a) < 2 or len(d) < 2 or len(m) < 2:
        raise ValueError("Los tres archivos deben contener encabezados y datos.")
    if len(a[0]) < 6 or len(d[0]) < 3 or "producto" not in normalize(a[0][1]) or "producto" not in normalize(d[0][0]):
        raise ValueError("Usa la estructura de los tres archivos originales: asignación con seis columnas, necesidad con tres y matriz con códigos en filas y columnas.")
    demands = {}
    for row in d[1:]:
        if all(x is None for x in row):
            continue
        pid = code(row[0])
        if pid in demands:
            raise ValueError(f"Necesidad duplicada para {pid}.")
        demands[pid] = row[2]
    products = []
    for row in a[1:]:
        if all(x is None for x in row):
            continue
        pid = code(row[1])
        if pid not in demands:
            raise ValueError(f"Falta necesidad para {pid}; registra cero si no se requiere.")
        products.append({"id": pid, "description": str(row[2]), "line": str(row[0]), "rate": row[5], "demand": demands[pid]})
    if set(demands) != {p["id"] for p in products}:
        raise ValueError("Hay productos de la necesidad sin asignación a una línea.")
    columns = [code(v) for v in m[0][1:] if v is not None]
    matrix = {}
    seen_rows = set()
    if len(columns) != len(set(columns)):
        raise ValueError("La matriz tiene encabezados duplicados.")
    for row in m[1:]:
        if row[0] is None:
            continue
        pid = code(row[0])
        if pid in seen_rows:
            raise ValueError("La matriz tiene filas duplicadas.")
        seen_rows.add(pid)
        for j, other in enumerate(columns, 1):
            v = row[j] if j < len(row) else None
            matrix[f"{pid}|{other}"] = 0 if pid == other and v is None else v
    return {"version": 1, "name": "Escenario importado", "products": products, "matrix": matrix, "config": dict(base_config, initial=None), "closures": []}


def excel_export(result):
    """Runtime download feature: create a standalone, formula-audited workbook."""
    wb = openpyxl.Workbook()
    cover = wb.active
    cover.title = "Resumen"
    cover.append(["Plan de fabricación", result["scenario"]["name"]])
    for row in [
        ["Fecha de cálculo", datetime.now().isoformat(timespec="seconds")],
        ["Estado del optimizador", result["optimization"]["status"]],
        ["Objetivo", "Minimizar minutos de cambio; no certifica mínimo plazo o costo total"],
        ["Necesidad (cartones)", result["demand"]], ["Programado", result["produced"]], ["Pendiente", result["missing"]],
        ["Cambios (min)", result["setup_minutes"]], ["Cota inferior (min)", result["optimization"]["bound"]],
        ["Capacidad neta (min)", result["available_minutes"]], ["Cierre", result["finish"] or "Sin cierre en el horizonte"],
        ["Huella del escenario", result["fingerprint"]],
        ["Supuestos", "Una línea; una campaña por producto; tasas por 8 h; cartones enteros; pausas reanudables; paradas al inicio del turno; sin fechas intermedias ni restricciones de materiales."],
    ]:
        cover.append(row)
    products = wb.create_sheet("Productos")
    products.append(["Código", "Descripción", "Línea", "Necesidad base", "Cartones por 8 horas"])
    for p in result["scenario"]["products"]:
        products.append([p["id"], p["description"], p["line"], p["demand"], p["rate"]])
    detail = wb.create_sheet("Detalle")
    detail.append(["Fecha", "Turno", "Código", "Descripción", "Cartones", "Min producción", "Min cambio"])
    for r in result["details"]:
        detail.append([r["date"], r["shift"], r["id"], r["description"], r["quantity"], r["production_minutes"], r["setup_minutes"]])
    shifts = wb.create_sheet("Turnos")
    shifts.append(["Fecha", "Turno", "Disponibles min", "Producción min", "Cambios min", "Ocupados min", "Libres min", "Control capacidad"])
    for i, r in enumerate(result["shifts"], 2):
        shifts.append([r["date"], r["shift"], r["available_minutes"], r["production_minutes"], r["setup_minutes"], f"=D{i}+E{i}", f"=C{i}-F{i}", f'=IF(ROUND(F{i},8)<=ROUND(C{i},8),"OK","SOBRECARGA")'])
    plan = wb.create_sheet("Plan por turno")
    keys = [(r["date"], r["shift"]) for r in result["shifts"]]
    plan.append(["Código", "Descripción"] + [f"{d} T{t}" for d, t in keys] + ["Total"])
    quantities = {}
    for r in result["details"]:
        k = (r["id"], r["date"], r["shift"])
        quantities[k] = quantities.get(k, 0) + r["quantity"]
    from openpyxl.utils import get_column_letter
    for i, pid in enumerate(result["route"], 2):
        p = next(p for p in result["scenario"]["products"] if p["id"] == pid)
        vals = [quantities.get((pid, d, t), 0) for d, t in keys]
        plan.append([pid, p["description"]] + vals + [f"=SUM(C{i}:{get_column_letter(2+len(keys))}{i})" if keys else "=0"])
    plan.freeze_panes = "C2"
    checks = wb.create_sheet("Cumplimiento")
    checks.append(["Código", "Necesidad escenario", "Programado", "Pendiente", "Control"])
    for i, pid in enumerate(result["made"], 2):
        checks.append([pid, result["made"][pid] + result["remaining"][pid], result["made"][pid], f"=B{i}-C{i}", f'=IF(D{i}=0,"COMPLETO","PENDIENTE")'])
    config = wb.create_sheet("Escenario")
    config.append(["Parámetro", "Valor"])
    for k, v in result["scenario"]["config"].items():
        config.append([k, v])
    matrix = wb.create_sheet("Cambios")
    ids = [p["id"] for p in result["scenario"]["products"]]
    matrix.append(["De / A"] + ids)
    for pid in ids:
        matrix.append([pid] + [result["scenario"]["matrix"].get(f"{pid}|{q}") for q in ids])
    for sheet in wb:
        sheet.freeze_panes = sheet.freeze_panes or "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="173B4B")
            cell.font = Font(color="FFFFFF", bold=True)
        for row in sheet:
            for c in row:
                formula_allowed = c.row > 1 and ((sheet.title == "Turnos" and c.column in (6, 7, 8)) or (sheet.title == "Plan por turno" and c.column == sheet.max_column) or (sheet.title == "Cumplimiento" and c.column in (4, 5)))
                if isinstance(c.value, str) and c.value.startswith("=") and not formula_allowed:
                    c.data_type = "s"  # User input remains literal text, including leading equals.
                c.alignment = Alignment(vertical="top", wrap_text=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = 23 if letter != "B" else 48
        sheet.row_dimensions[1].height = 32
    cover.column_dimensions["B"].width = 95
    cover.row_dimensions[13].height = 60
    for c in range(3, plan.max_column + 1):
        plan.column_dimensions[get_column_letter(c)].width = 17
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
